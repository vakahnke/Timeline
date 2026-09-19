from datetime import timedelta
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.utils import timezone
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches
from rest_framework.test import APIClient

from projects.models import Project, ProjectMembership, Role

from .models import Category, Event
from .pptx_export import build_pptx
from .status_report import build_facts, suggest

User = get_user_model()
PPTX = 'application/vnd.openxmlformats-officedocument.presentationml.presentation'


def _walk(shapes):
    for s in shapes:
        yield s
        if s.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _walk(s.shapes)


def _all_text(slide):
    out = []
    for s in _walk(slide.shapes):
        if s.has_text_frame:
            out.append(s.text_frame.text)
        if getattr(s, 'has_table', False) and s.has_table:
            out += [c.text_frame.text for r in s.table.rows for c in r.cells]
    return '\n'.join(out)


class _Base(TestCase):
    def setUp(self):
        self.owner = User.objects.create_user('pm', 'pm@example.com', 'a-real-Passw0rd')
        self.viewer = User.objects.create_user('lead', 'lead@example.com', 'a-real-Passw0rd')
        self.outsider = User.objects.create_user('nobody', 'n@example.com', 'a-real-Passw0rd')
        self.project = Project.objects.create(name='Launch & <Growth>', owner=self.owner)
        ProjectMembership.objects.create(project=self.project, user=self.owner, role=Role.OWNER)
        ProjectMembership.objects.create(project=self.project, user=self.viewer, role=Role.VIEWER)
        for name, color in (('Build', '#2f6fe0'), ('Launch', '#239b73')):
            Category.objects.create(project=self.project, name=name, color=color)
        now = timezone.now().replace(hour=12, minute=0, second=0, microsecond=0)
        a = Event.objects.create(project=self.project, title='Scope locked', category='Build', is_milestone=True,
                                 start=now - timedelta(days=20), end=now - timedelta(days=10), percent_complete=100)
        b = Event.objects.create(project=self.project, title='Build the core flow', category='Build', is_milestone=True,
                                 start=now - timedelta(days=10), end=now + timedelta(days=20), percent_complete=35)
        c = Event.objects.create(project=self.project, title='Public launch', category='Launch', is_milestone=True,
                                 start=now + timedelta(days=20), end=now + timedelta(days=30), percent_complete=0)
        b.depends_on.set([a])
        c.depends_on.set([b])
        self.project.committed_end = timezone.localtime(now + timedelta(days=27)).date()
        self.project.save()
        import json
        self.facts = json.loads(json.dumps(build_facts(self.project, now=now)))   # exactly what the API sends
        s = suggest(self.facts)
        self.report = {'status': s['status'], 'status_source': 'rule', 'override_reason': '', 'rule_fired': s['rule_fired']}
        self.doc = {
            'header': {'project': self.project.name, 'subtitle': 'Status report', 'date': 'Sep 19, 2026', 'pm': 'PM: pm'},
            'headline': 'Launch is three days late <unless> we cut scope & decide by Friday.',
            'pathToGreen': 'Cut the settings page from beta scope.',
            'show': {}, 'timeline': {},
            'decision': {'none': False, 'title': 'Decision needed', 'neededBy': 'Sep 26', 'from': 'the sponsor', 'text': 'Shorten the beta?'},
            'kpis': [{'id': 'finish', 'label': 'Forecast finish', 'value': 'Nov 7', 'detail': '+3 days', 'tone': 'warn'},
                     {'id': 'work', 'label': 'Work complete', 'value': '41%', 'detail': '42% elapsed', 'tone': ''}],
            'columns': [
                {'id': 'done', 'kind': 'list', 'mark': '✓', 'title': 'Last two weeks', 'items': [{'id': 'l1', 'text': 'Scope locked', 'when': 'Sep 9'}, {'id': 'l2', 'text': '', 'when': ''}]},
                {'id': 'risks', 'kind': 'risks', 'title': 'Top risks', 'items': [{'id': 'r1', 'severity': 'high', 'text': 'Core flow is behind.', 'detail': 'E. Editor · cut scope'}]},
                {'id': 'free', 'kind': 'text', 'title': 'Budget', 'text': 'On budget: $41k of $120k spent.'},
            ],
            'footer': {'text': ''},
        }

    def build(self, **over):
        kw = dict(doc=self.doc, report=self.report, facts=self.facts, layout='slide', paper='letter', previous=None, tz_offset_minutes=-240)
        kw.update(over)
        return Presentation(BytesIO(build_pptx(**kw)))


class PptxContentTests(_Base):
    def test_slide_is_16_9_and_every_element_is_native(self):
        prs = self.build()
        self.assertEqual((prs.slide_width, prs.slide_height), (Inches(13.333), Inches(7.5)))
        self.assertEqual(len(prs.slides), 1)
        shapes = list(_walk(prs.slides[0].shapes))
        kinds = {s.shape_type for s in shapes}
        self.assertNotIn(MSO_SHAPE_TYPE.PICTURE, kinds)                       # nothing is a screenshot
        self.assertTrue(kinds <= {MSO_SHAPE_TYPE.TEXT_BOX, MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.GROUP}, kinds)
        groups = [s for s in prs.slides[0].shapes if s.shape_type == MSO_SHAPE_TYPE.GROUP]
        self.assertEqual([g.name for g in groups], ['Timeline'])              # one group: moves as a unit, can be ungrouped

    def test_text_is_editable_and_markup_characters_survive(self):
        text = _all_text(self.build().slides[0])
        self.assertIn('Launch is three days late <unless> we cut scope & decide by Friday.', text)
        self.assertIn('Launch & <Growth>', text)
        self.assertIn('Path to green: Cut the settings page from beta scope.', text)
        self.assertIn('DECISION NEEDED · BY SEP 26 · FROM THE SPONSOR', text)
        self.assertIn('On budget: $41k of $120k spent.', text)
        self.assertIn('AT RISK', text.upper())
        self.assertNotIn('\n✓  \n', text)                                     # the empty list item is dropped

    def test_timeline_has_one_diamond_per_milestone_with_alt_text(self):
        slide = self.build().slides[0]
        grp = next(s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.GROUP)
        ms = [s for s in grp.shapes if s.name.startswith('Milestone: ')]
        self.assertEqual(sorted(s.name for s in ms), ['Milestone: Build the core flow', 'Milestone: Public launch', 'Milestone: Scope locked'])
        self.assertIn('diamond', str(ms[0].auto_shape_type).lower())
        nv = grp._element.xpath('.//p:nvGrpSpPr/p:cNvPr')[0]
        self.assertIn('3 milestones', nv.get('descr'))
        self.assertTrue(any(s.name == 'Today line' for s in grp.shapes))
        self.assertTrue(any(s.name == 'Critical path' for s in grp.shapes))

    def test_timeline_stays_inside_the_slide_and_bars_follow_the_dates(self):
        prs = self.build()
        grp = next(s for s in prs.slides[0].shapes if s.shape_type == MSO_SHAPE_TYPE.GROUP)
        for s in grp.shapes:
            self.assertGreaterEqual(s.left, 0, s.name)
            self.assertLessEqual(s.left + s.width, prs.slide_width + Inches(0.02), s.name)
            self.assertLessEqual(s.top + s.height, prs.slide_height, s.name)
        build = next(s for s in grp.shapes if s.name == 'Build planned')
        launch = next(s for s in grp.shapes if s.name == 'Launch planned')
        self.assertLess(build.left, launch.left)                              # Build starts first
        self.assertGreater(build.width, launch.width)                         # 40 days vs 10 days
        self.assertAlmostEqual(build.width / launch.width, 4.0, delta=0.15)   # …to scale

    def test_milestone_labels_never_cover_a_percent_figure(self):
        import json
        now = timezone.now().replace(hour=12, minute=0, second=0, microsecond=0)
        # A milestone in the row below, right under the Build row's "51%" figure.
        Event.objects.create(project=self.project, title='Beta invite list ready', category='Launch', is_milestone=True,
                             start=now - timedelta(days=6), end=now - timedelta(days=2), percent_complete=100)
        for i, name in enumerate(('Marketing', 'Sales', 'Support')):      # more tracks: rows get tight
            Event.objects.create(project=self.project, title=f'{name} work', category=name,
                                 start=now + timedelta(days=i), end=now + timedelta(days=25), percent_complete=0)
        facts = json.loads(json.dumps(build_facts(self.project, now=now)))
        for layout in ('slide', 'handout'):
            grp = next(s for s in self.build(facts=facts, layout=layout).slides[0].shapes if s.shape_type == MSO_SHAPE_TYPE.GROUP)
            figures = [s for s in grp.shapes if s.name == 'Percent complete']
            labels = [s for s in grp.shapes if s.name.startswith('Label: ')]
            self.assertTrue(figures and labels)
            for f in figures:
                mid = f.top + f.height // 2
                for lb in labels:
                    covers = lb.left < f.left + f.width and lb.left + lb.width > f.left and lb.top < mid < lb.top + lb.height
                    self.assertFalse(covers, f'{layout}: "{lb.name}" covers a percent figure')

    def test_hidden_blocks_and_tracks_are_left_out(self):
        doc = {**self.doc, 'show': {'kpis': False, 'timeline': False, 'footer': False}}
        slide = self.build(doc=doc).slides[0]
        self.assertFalse([s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.GROUP])
        text = _all_text(slide)
        self.assertNotIn('FORECAST FINISH', text)
        self.assertNotIn('Schedule data as of', text)
        doc = {**self.doc, 'timeline': {'hiddenRows': ['Launch'], 'milestoneIds': None}}
        grp = next(s for s in self.build(doc=doc).slides[0].shapes if s.shape_type == MSO_SHAPE_TYPE.GROUP)
        names = [s.name for s in grp.shapes]
        self.assertNotIn('Track: Launch', names)
        self.assertNotIn('Milestone: Public launch', names)                   # its track is hidden

    def test_override_and_no_decision_are_stated(self):
        rep = {**self.report, 'status': 'on_track', 'status_source': 'override', 'override_reason': 'Sponsor moved the date on Sep 18.'}
        doc = {**self.doc, 'decision': {'none': True}}
        text = _all_text(self.build(report=rep, doc=doc, previous={'status': 'at_risk', 'as_of': self.facts['as_of']}).slides[0])
        self.assertIn('status set by the author: Sponsor moved the date on Sep 18.', text)
        self.assertIn('NO DECISIONS NEEDED', text)
        self.assertIn('was At risk', text)                                    # trend against the previous report

    def test_handout_is_portrait_with_a_real_table(self):
        prs = self.build(layout='handout', paper='a4')
        self.assertLess(prs.slide_width, prs.slide_height)
        self.assertAlmostEqual(prs.slide_width / 914400, 8.268, places=2)
        tables = [s for s in prs.slides[0].shapes if getattr(s, 'has_table', False) and s.has_table]
        self.assertEqual(len(tables), 1)
        t = tables[0].table
        self.assertEqual([c.text_frame.text for c in t.rows[0].cells], ['MILESTONE', 'TRACK', 'DATE', 'DONE', 'STATUS'])
        self.assertEqual(len(t.rows), 4)
        self.assertIn('Met', t.cell(1, 4).text_frame.text)

    def test_empty_project_still_produces_a_valid_file(self):
        empty = Project.objects.create(name='Empty', owner=self.owner)
        import json
        facts = json.loads(json.dumps(build_facts(empty)))
        prs = self.build(facts=facts, doc={**self.doc, 'kpis': [], 'columns': []})
        self.assertEqual(len(prs.slides), 1)

    def test_provenance_travels_in_the_speaker_notes(self):
        notes = self.build().slides[0].notes_slide.notes_text_frame.text
        self.assertIn('Schedule data as of', notes)
        self.assertIn('native, editable', notes)


class PptxEndpointTests(_Base):
    def url(self):
        return f'/api/projects/{self.project.id}/status-reports/export-pptx/'

    def payload(self, **over):
        p = {'layout': 'slide', 'paper': 'letter', **self.report, 'content': self.doc, 'snapshot': self.facts, 'previous': None, 'tz_offset': -240}
        p.update(over)
        return p

    def client_for(self, user):
        c = APIClient()
        c.force_authenticate(user)
        return c

    def test_owner_and_viewer_can_download_outsider_cannot(self):
        for user in (self.owner, self.viewer):
            r = self.client_for(user).post(self.url(), self.payload(), format='json')
            self.assertEqual(r.status_code, 200, getattr(r, 'data', None))
            self.assertEqual(r['Content-Type'], PPTX)
            self.assertIn('launch-growth-status-', r['Content-Disposition'])
            self.assertEqual(len(Presentation(BytesIO(r.content)).slides), 1)
        self.assertEqual(self.client_for(self.outsider).post(self.url(), self.payload(), format='json').status_code, 403)
        self.assertEqual(APIClient().post(self.url(), self.payload(), format='json').status_code, 401)

    def test_bad_input_is_rejected(self):
        c = self.client_for(self.owner)
        self.assertEqual(c.post(self.url(), self.payload(status='great'), format='json').status_code, 400)
        self.assertEqual(c.post(self.url(), self.payload(content=['no']), format='json').status_code, 400)
        self.assertEqual(c.post(self.url(), self.payload(snapshot={'no': 'as_of'}), format='json').status_code, 400)
        self.assertEqual(c.post(self.url(), self.payload(content={'x': 'y' * 250_000}), format='json').status_code, 400)
        self.assertEqual(c.post(self.url(), self.payload(tz_offset=99999), format='json').status_code, 400)

    def test_exporting_saves_nothing(self):
        from .models import StatusReport
        self.client_for(self.viewer).post(self.url(), self.payload(), format='json')
        self.assertEqual(StatusReport.objects.count(), 0)

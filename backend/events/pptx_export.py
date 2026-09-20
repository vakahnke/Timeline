"""Native PowerPoint export of a status report.

Everything on the slide is a real PowerPoint object: text boxes you can retype, rectangles and
diamonds for the timeline (grouped, so it moves as one and can be ungrouped), a real table on the
handout. Nothing is a picture. The input is exactly what the print tool holds: the report document
(``content``), the status fields, and the facts it was drawn from, so the file matches the page on
screen, including edits that have not been saved.

Design notes (docs/design/status-one-pager.md §2 "Export technology", §3.5):
- Theme fonts (``+mj-lt`` / ``+mn-lt``) so the host deck's typography applies on paste, fixed
  sizes with wrapping, **no autofit** (only PowerPoint recomputes it), and slack in every box
  because Keynote and Google Slides substitute fonts.
- Status is a glyph + a word + a colour, in one text run, so it survives editing and grey-scale.
- Shapes are added in reading order, which is the accessibility reading order.
- Geometry mirrors ``frontend/src/report.css`` (1cqw = 0.1333in on the slide). The golden test in
  ``tests_pptx_export.py`` pins it.
"""
import math
from datetime import datetime, timedelta, timezone as dt_tz
from io import BytesIO

from django.conf import settings
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

INK, INK2, INK3, RULE = '16202E', '4B5768', '7C8797', 'D9DEE6'
ACCENT, OK, WARN, BAD = '2B50C8', '1C7C54', 'B86E00', 'B3362B'
TONE = {'ok': OK, 'warn': WARN, 'bad': BAD, 'ink': INK, '': INK}
STATUS = {
    'on_track':  {'label': 'On track',  'glyph': '●', 'color': OK,   'fill': 'E6F4EC'},
    'at_risk':   {'label': 'At risk',   'glyph': '◆', 'color': WARN, 'fill': 'FFF4E0'},
    'off_track': {'label': 'Off track', 'glyph': '■', 'color': BAD,  'fill': 'FBE9E7'},
}
ORDER = ['on_track', 'at_risk', 'off_track']
SEVERITY = {'high': ('■', BAD), 'medium': ('◆', WARN), 'low': ('●', OK)}
MONTHS = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec']
HEAD_FONT, BODY_FONT = '+mj-lt', '+mn-lt'
PAPER = {'letter': (8.5, 11.0), 'a4': (8.268, 11.693)}


# ── small drawing helpers (all coordinates in inches) ───────────────────────────────────────────
def _rgb(h):
    return RGBColor.from_string(h.lstrip('#').upper()[:6].ljust(6, '0'))


def _describe(shape, name, alt=None):
    shape.name = name
    if alt:
        for el in shape._element.iter(qn('p:cNvPr')):
            el.set('descr', alt)
            break
    return shape


def _box(shapes, x, y, w, h, fill=None, line=None, line_pt=0.75, kind=MSO_SHAPE.RECTANGLE, radius=None, name='Shape'):
    s = shapes.add_shape(kind, Inches(x), Inches(y), Inches(max(w, 0.005)), Inches(max(h, 0.005)))
    s.shadow.inherit = False
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = _rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = _rgb(line)
        s.line.width = Pt(line_pt)
    else:
        s.line.fill.background()
    if radius is not None and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.name = name
    return s


def _frame(shape, anchor=MSO_ANCHOR.TOP, inset=0.0):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(inset)
    return tf


def _write(tf, paragraphs, size, color=INK, font=BODY_FONT, align=PP_ALIGN.LEFT, space_after=0, line_spacing=None):
    """``paragraphs``: list of paragraphs; each a string or a list of runs ``(text, opts)`` where
    opts may set bold / color / size / italic / font."""
    first = True
    for para in paragraphs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_after = Pt(space_after)
        if line_spacing:
            p.line_spacing = line_spacing
        runs = [(para, {})] if isinstance(para, str) else para
        for text, o in runs:
            r = p.add_run()
            r.text = text
            f = r.font
            f.size = Pt(o.get('size', size))
            f.bold = bool(o.get('bold'))
            f.italic = bool(o.get('italic'))
            f.name = o.get('font', font)
            f.color.rgb = _rgb(o.get('color', color))


def _text(shapes, x, y, w, h, paragraphs, size, name='Text', anchor=MSO_ANCHOR.TOP, **kw):
    tb = shapes.add_textbox(Inches(x), Inches(y), Inches(max(w, 0.05)), Inches(max(h, 0.05)))
    tb.name = name
    _write(_frame(tb, anchor), paragraphs, size, **kw)
    return tb


def _est_width(text, pt, bold=False):
    """Rough rendered width in inches; generous, because fonts are substituted across apps."""
    return len(text) * pt * (0.58 if bold else 0.54) / 72


def _lines(text, pt, width, bold=False):
    """Estimated number of wrapped lines for ``text`` in a box ``width`` inches wide."""
    # Typical widths (about 0.47em a character), not the generous figure _est_width uses to keep
    # labels from colliding: reserving space for lines that never appear starves the timeline.
    if not text:
        return 0
    return max(1, math.ceil(len(text) * pt * (0.5 if bold else 0.47) / 72 / max(0.5, width)))


def _fmt_day(dt):
    return f'{MONTHS[dt.month - 1]} {dt.day}'


def _parse(iso, tz):
    return datetime.fromisoformat(iso.replace('Z', '+00:00')).astimezone(tz)


# ── report pieces ───────────────────────────────────────────────────────────────────────────────
def _chosen_milestones(doc, facts, tz):
    events = facts.get('events') or []
    ids = (doc.get('timeline') or {}).get('milestoneIds')
    picked = [e for e in events if e['id'] in ids] if ids is not None else [e for e in events if e.get('is_milestone')]
    now = _parse(facts['as_of'], tz)
    out = []
    for e in picked:
        date = _parse(e['end'], tz)
        state = 'done' if e.get('percent_complete', 0) >= 100 else ('late' if date < now else 'upcoming')
        out.append({**e, 'date': date, 'state': state, 'baseline': _parse(e['baseline_end'], tz) if e.get('baseline_end') else None})
    return sorted(out, key=lambda m: m['date'])


def _trend(report, previous, tz):
    if not previous or previous.get('status') not in STATUS:
        return ''
    when = _fmt_day(_parse(previous['as_of'], tz)) if previous.get('as_of') else ''
    if previous['status'] == report['status']:
        return f'  ●  unchanged since {when}'
    worse = ORDER.index(report['status']) > ORDER.index(previous['status'])
    return f"  {'▼' if worse else '▲'}  was {STATUS[previous['status']]['label']} · {when}"


def _footer_text(doc, report, facts, tz):
    custom = ((doc.get('footer') or {}).get('text') or '').strip()
    if custom:
        return custom
    as_of = _parse(facts['as_of'], tz)
    parts = [f'Schedule data as of {_fmt_day(as_of)}, {as_of.year}']
    committed = (facts.get('project') or {}).get('committed_end')
    if committed:
        y, m, d = (int(v) for v in committed.split('-'))
        word = 'baseline' if (facts.get('project') or {}).get('commitment_source') == 'baseline' else 'committed'
        parts.append(f'{word} finish {MONTHS[m - 1]} {d}')
    else:
        parts.append('no committed finish date set')
    if facts.get('baseline') and (doc.get('show') or {}).get('baseline') is True:
        b = facts['baseline']
        parts.append(f"baseline “{b.get('name') or ''}” of {_fmt_day(_parse(b['created_at'], tz))}")
    if report.get('status_source') == 'override':
        parts.append(f"status set by the author: {report.get('override_reason') or 'no reason given'}")
    elif report.get('rule_fired'):
        parts.append(f"{STATUS[report['status']]['label']} by rule: {report['rule_fired']}")
    return ' · '.join(parts)


def _signed(days):
    return f"{'+' if days > 0 else '−'}{abs(days)}d"


def _slip_text(days):
    return '–' if days is None else 'on plan' if days == 0 else _signed(days)


def _draw_timeline(slide, x0, y0, w, h, doc, facts, tz, pt):
    """The simplified timeline as grouped native shapes. Mirrors ReportTimeline.jsx."""
    tl = doc.get('timeline') or {}
    hidden = set(tl.get('hiddenRows') or [])
    rows = [r for r in (facts.get('rows') or []) if r['name'] not in hidden]
    if facts.get('empty') or not rows or h < 0.6:
        return None
    milestones = _chosen_milestones(doc, facts, tz)
    show_progress, show_critical = tl.get('showProgress', True), tl.get('showCritical', True)
    show_baseline = bool(facts.get('baseline')) and (doc.get('show') or {}).get('baseline') is True      # opt-in, per report
    one_day = timedelta(days=1)

    def moved(a, b):
        return bool(a and b) and abs(a - b) >= one_day

    # Baseline ghosts only where the plan moved by a day or more, so an on-plan chart stays clean.
    ghosts = {}
    if show_baseline:
        for r in rows:
            if r.get('baseline_start') and r.get('baseline_end'):
                bs, be = _parse(r['baseline_start'], tz), _parse(r['baseline_end'], tz)
                if moved(bs, _parse(r['start'], tz)) or moved(be, _parse(r['end'], tz)):
                    ghosts[r['name']] = (bs, be)
    ghost_ms = {m['id'] for m in milestones if show_baseline and moved(m.get('baseline'), m['date'])}
    base_times = [t for pair in ghosts.values() for t in pair] + [m['baseline'] for m in milestones if m['id'] in ghost_ms]

    grp = slide.shapes.add_group_shape()
    g = grp.shapes
    u = pt / 72                                   # one text line-height unit, in inches
    label_w = min(w * 0.17, max(1.15, max(_est_width(r['name'][:17], pt, True) for r in rows) + u * 1.2))
    right_pad, top, legend_h = u * 1.2, u * 3.0, u * 2.2
    body_h = h - top - legend_h
    row_h = body_h / len(rows)
    bar_h = min(row_h * 0.42, u * 1.7)
    dia = u * 0.62

    start, end = _parse(facts['start'], tz), _parse(facts['end'], tz)
    t0 = min([start] + base_times) - timedelta(days=4)
    t1 = max([end] + [m['date'] for m in milestones] + base_times) + timedelta(days=7)
    span = (t1 - t0).total_seconds()
    px0, px1 = x0 + label_w, x0 + w - right_pad

    def X(t):
        return px0 + (t - t0).total_seconds() / span * (px1 - px0)

    today = _parse(facts['as_of'], tz)

    # month bands + labels
    c = t0.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    i = 0
    while c < t1 and i < 60:
        nxt = (c.replace(day=28) + timedelta(days=4)).replace(day=1)
        a, b = max(t0, c), min(t1, nxt)
        if b > a:
            if i % 2 == 1:
                _box(g, X(a), y0 + top, X(b) - X(a), body_h, fill='F6F7FA', name='Month band')
            if X(b) - X(a) > u * 2.4:
                _text(g, X(a) + u * 0.3, y0, max(0.5, X(b) - X(a)), u * 1.5, [[(MONTHS[c.month - 1].upper(), {'bold': True})]], pt * 0.9, color=INK3, name='Month')
        c, i = nxt, i + 1
    if today > t0:
        past = _box(g, px0, y0 + top, min(X(today), px1) - px0, body_h, fill='E9EBEF', name='Past')
        past.fill.fore_color.rgb = _rgb('EEF0F3')
    _box(g, px0, y0 + top - 0.005, px1 - px0, 0.01, fill=RULE, name='Axis')

    obstacles = []                             # percent texts and diamonds: labels must not cover them
    for idx, r in enumerate(rows):
        cy = y0 + top + row_h * idx + row_h / 2
        rs, re = _parse(r['start'], tz), _parse(r['end'], tz)
        xs, xe = X(rs), X(re)
        bw = max(0.03, xe - xs)
        color = (r.get('color') or '#6B7A90').lstrip('#')
        if idx:
            _box(g, x0, y0 + top + row_h * idx - 0.004, w - right_pad, 0.008, fill='EDF0F4', name='Row rule')
        _box(g, x0, cy - bar_h / 2, u * 0.36, bar_h, fill=color, name='Track colour')
        name = r['name'] if len(r['name']) <= 17 else r['name'][:16] + '…'
        _text(g, x0 + u * 0.8, cy - u * 0.8, label_w - u * 0.9, u * 1.6, [[(name, {'bold': True})]], pt * 1.04, anchor=MSO_ANCHOR.MIDDLE, name=f'Track: {r["name"]}')
        light = _box(g, xs, cy - bar_h / 2, bw, bar_h, fill=color, line=color, line_pt=0.75, kind=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12, name=f'{r["name"]} planned')
        _lighten(light, 0.78)
        done = max(0.0, min(1.0, (r.get('progress') or 0) / 100))
        if show_progress and done > 0:
            _box(g, xs, cy - bar_h / 2, max(0.02, bw * done), bar_h, fill=color, kind=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12, name=f'{r["name"]} done')
            if done < 1 and bw * (1 - done) > u * 3:
                _text(g, xs + bw * done + u * 0.35, cy - u * 0.75, u * 3.2, u * 1.5, [[(f'{round(done * 100)}%', {'bold': True})]], pt * 0.9, anchor=MSO_ANCHOR.MIDDLE, name='Percent complete')
                obstacles.append((xs + bw * done + u * 0.35, cy - u * 0.6, xs + bw * done + u * 3.0, cy + u * 0.6))
        if r['name'] in ghosts:                    # the baseline: a thin outline just above the bar
            bs, be = ghosts[r['name']]
            gh = max(0.03, min(row_h * 0.13, u * 0.42))
            _describe(_box(g, X(bs), cy - bar_h / 2 - u * 0.22 - gh, max(0.03, X(be) - X(bs)), gh, fill='FFFFFF', line=INK3, line_pt=0.75, name=f'{r["name"]} baseline'),
                      f'{r["name"]} baseline', f'{r["name"]} baseline: {_fmt_day(bs)} to {_fmt_day(be)}')
        if show_critical:
            for a, b in r.get('critical_spans') or []:
                xa, xb = X(_parse(a, tz)), X(_parse(b, tz))
                _box(g, xa, cy + bar_h / 2 + u * 0.25, max(0.03, xb - xa), max(0.022, u * 0.17), fill=ACCENT, name='Critical path')

    if t0 < today < t1:
        _box(g, X(today) - 0.009, y0 + top - 0.03, 0.018, body_h + 0.03, fill=BAD, name='Today line')
        flag = _box(g, X(today) - u * 2.1, y0, u * 4.2, u * 1.5, fill=BAD, kind=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.18, name='Today')
        _write(_frame(flag, MSO_ANCHOR.MIDDLE), [[('TODAY', {'bold': True})]], pt * 0.82, color='FFFFFF', align=PP_ALIGN.CENTER)

    # milestones: diamonds, then labels placed without collisions (same search as the client)
    row_index = {r['name']: i for i, r in enumerate(rows)}
    placed = list(obstacles)
    for m in milestones:
        ri = row_index.get(m.get('category'), row_index.get('Other'))
        if ri is not None:
            mx, cy = X(m['date']), y0 + top + row_h * ri + row_h / 2
            # bar-height box, so a diamond never blocks its own label just above or below it
            placed.append((mx - dia, cy - bar_h / 2, mx + dia, cy + bar_h / 2))
            if m['id'] in ghost_ms:
                bx_ = X(m['baseline'])
                placed.append((bx_ - dia * 0.7, cy - bar_h / 2, bx_ + dia * 0.7, cy + bar_h / 2))

    def clear(bx):
        if bx[0] < px0 - u or bx[2] > x0 + w - 0.02 or bx[1] < y0 + u * 1.5 or bx[3] > y0 + h - legend_h + u * 0.4:
            return False
        return not any(bx[0] < o[2] + u * 0.4 and bx[2] > o[0] - u * 0.4 and bx[1] < o[3] and bx[3] > o[1] for o in placed)

    for m in milestones:
        ri = row_index.get(m.get('category'), row_index.get('Other'))
        if ri is None:
            continue
        mx, cy = X(m['date']), y0 + top + row_h * ri + row_h / 2
        if m['id'] in ghost_ms:                     # hollow baseline diamond, joined to where it is now
            bx_, q = X(m['baseline']), dia * 0.7
            _box(g, min(bx_, mx), cy - 0.006, max(0.02, abs(mx - bx_)), 0.012, fill=INK3, name='Slip')
            _describe(_box(g, bx_ - q, cy - q, q * 2, q * 2, fill='FFFFFF', line=INK3, line_pt=0.75, kind=MSO_SHAPE.DIAMOND, name=f'Baseline: {m["title"]}'),
                      f'Baseline: {m["title"]}', f'{m["title"]} baseline date {_fmt_day(m["baseline"])}')
        fill = INK if m['state'] == 'done' else BAD if m['state'] == 'late' else 'FFFFFF'
        line = INK if m['state'] == 'upcoming' else 'FFFFFF'
        _describe(_box(g, mx - dia, cy - dia, dia * 2, dia * 2, fill=fill, line=line, line_pt=1, kind=MSO_SHAPE.DIAMOND, name=f'Milestone: {m["title"]}'),
                  f'Milestone: {m["title"]}', f'{m["title"]}, {_fmt_day(m["date"])}, {m["state"]}')
        title = m['title'] if len(m['title']) <= 34 else m['title'][:33] + '…'
        short = m['title'] if len(m['title']) <= 16 else m['title'][:15] + '…'
        day = _fmt_day(m['date'])
        chosen = None
        slip = f" ({_signed(m['slip_days'])})" if m['id'] in ghost_ms and m.get('slip_days') else ''
        for text in (f'{title} · {day}{slip}', f'{short} · {day}{slip}', f'{day}{slip}', day):
            tw = _est_width(text, pt * 0.95, True)
            # Normally to the right of the diamond; go left first when a neighbour in the same row
            # sits within this label's reach, so that neighbour keeps a spot of its own.
            crowded = any(o is not m and row_index.get(o.get('category'), row_index.get('Other')) == ri
                          and 0 < X(o['date']) - mx < tw + dia * 2 + u for o in milestones)
            spots = ((True, True), (True, False), (False, True), (False, False)) if crowded else ((True, False), (True, True), (False, False), (False, True))
            for above, right in spots:
                ty = cy - bar_h / 2 - u * 1.45 if above else cy + bar_h / 2 + u * 0.5
                tx = mx - dia - u * 0.3 - tw if right else mx + dia + u * 0.3
                bx = (tx, ty, tx + tw, ty + u * 1.2)
                if clear(bx):
                    chosen = (text, tx, ty, tw, right)
                    break
            if chosen:
                break
        if chosen:
            text, tx, ty, tw, right = chosen
            placed.append((tx, ty, tx + tw, ty + u * 1.2))
            lbl = _text(g, tx, ty, tw + u * 0.6, u * 1.35, [[(text, {'bold': True})]], pt * 0.95, color=BAD if m['state'] == 'late' else INK,
                        align=PP_ALIGN.RIGHT if right else PP_ALIGN.LEFT, name=f'Label: {m["title"]}')
            if right:                              # the box was widened for slack: keep its right edge on the anchor
                lbl.left = Inches(tx - u * 0.6)

    # legend
    items = ([('line', 'Critical path')] if show_critical else []) + [('done', 'Milestone met'), ('up', 'Milestone ahead'), ('late', 'Past due')] + ([('base', 'Baseline')] if base_times else [])
    step = u * 9.4
    lx = x0 + w - right_pad - len(items) * step
    ly = y0 + h - u * 1.5
    for k, (kind, label) in enumerate(items):
        cx = lx + k * step
        if kind == 'line':
            _box(g, cx, ly + u * 0.62, u * 1.6, 0.025, fill=ACCENT, name='Legend')
        elif kind == 'base':
            _box(g, cx, ly + u * 0.5, u * 1.6, u * 0.4, fill='FFFFFF', line=INK3, line_pt=0.75, name='Legend')
        else:
            q = u * 0.45
            _box(g, cx + u * 0.3, ly + u * 0.7 - q, q * 2, q * 2, fill=INK if kind == 'done' else BAD if kind == 'late' else 'FFFFFF',
                 line=INK if kind == 'up' else 'FFFFFF', line_pt=0.75, kind=MSO_SHAPE.DIAMOND, name='Legend')
        _text(g, cx + u * 1.9, ly, step - u * 1.9, u * 1.4, [label], pt * 0.86, color=INK3, anchor=MSO_ANCHOR.MIDDLE, name='Legend')

    first, last = _fmt_day(start), _fmt_day(end)
    _describe(grp, 'Timeline', f'Timeline from {first} to {last}: {len(rows)} tracks, {len(milestones)} milestones. Grouped native shapes; ungroup to edit.')
    return grp


def _lighten(shape, amount):
    """Tint a solid fill toward white (python-pptx has no alpha API; a tint prints reliably)."""
    rgb = shape.fill.fore_color.rgb
    mixed = tuple(round(c + (255 - c) * amount) for c in (rgb[0], rgb[1], rgb[2]))
    shape.fill.fore_color.rgb = RGBColor(*mixed)


def _kpis(slide, x, y, w, h, kpis, sizes, per_row):
    rows = math.ceil(len(kpis) / per_row)
    cell_w, cell_h = w / per_row, h / rows
    _box(slide.shapes, x, y, w, 0.012, fill=RULE, name='Rule')
    _box(slide.shapes, x, y + h - 0.012, w, 0.012, fill=RULE, name='Rule')
    for i, k in enumerate(kpis):
        r, c = divmod(i, per_row)
        cx, cy = x + c * cell_w, y + r * cell_h
        if c:
            _box(slide.shapes, cx, cy + 0.04, 0.012, cell_h - 0.08, fill=RULE, name='Rule')
        if r and c == 0:
            _box(slide.shapes, x, cy, w, 0.012, fill=RULE, name='Rule')
        pad = 0.0 if c == 0 else 0.14
        _text(slide.shapes, cx + pad, cy + 0.07, cell_w - pad - 0.08, cell_h - 0.1, [
            [((k.get('label') or '').upper(), {'bold': True, 'color': INK3, 'size': sizes[0]})],
            [(k.get('value') or '', {'bold': True, 'color': TONE.get(k.get('tone') or '', INK), 'size': sizes[1], 'font': HEAD_FONT})],
            [(k.get('detail') or '', {'color': INK2, 'size': sizes[2]})],
        ], sizes[2], name=f"Number: {k.get('label') or ''}")


def _block(slide, x, y, w, h, col, sizes):
    """One text block (list / risks / free text): a heading, a rule, and ONE text box holding all
    the items as paragraphs, so rewording reflows naturally."""
    _text(slide.shapes, x, y, w, 0.24, [[((col.get('title') or '').upper(), {'bold': True})]], sizes[0], color=INK3, name='Block heading')
    _box(slide.shapes, x, y + 0.27, w, 0.012, fill=RULE, name='Rule')
    kind, paras = col.get('kind'), []
    if kind == 'text':
        paras = [col.get('text') or '']
    elif kind == 'risks':
        for it in [i for i in col.get('items') or [] if (i.get('text') or i.get('detail'))]:
            glyph, color = SEVERITY.get(it.get('severity'), SEVERITY['medium'])
            runs = [(glyph + '  ', {'color': color, 'size': sizes[1] * 0.8}), ((it.get('text') or '') + ' ', {'bold': True})]
            if it.get('detail'):
                runs.append((it['detail'], {'color': INK2}))
            paras.append(runs)
        paras = paras or [[('No significant risks', {'italic': True, 'color': INK3})]]
    else:
        for it in [i for i in col.get('items') or [] if i.get('text')]:
            runs = [((col.get('mark') or '•') + '  ', {'bold': True, 'color': INK3}), (it['text'], {})]
            if it.get('when'):
                runs.append((f" · {it['when']}", {'color': INK3}))
            paras.append(runs)
        paras = paras or [[('Nothing to report', {'italic': True, 'color': INK3})]]
    _text(slide.shapes, x, y + 0.34, w, max(0.3, h - 0.34), paras, sizes[1], space_after=sizes[1] * 0.28, name=col.get('title') or 'Block')


def _ask(slide, x, y, w, h, doc, size):
    d = doc.get('decision') or {}
    if d.get('none'):
        box = _box(slide.shapes, x, y, w, h, fill='F7F8FA', line=RULE, line_pt=1, kind=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06, name='Decision')
        _write(_frame(box, MSO_ANCHOR.MIDDLE, 0.1), [[('NO DECISIONS NEEDED', {'bold': True})]], size * 0.86, color=INK3, align=PP_ALIGN.CENTER)
        return
    head = (d.get('title') or 'Decision needed')
    if d.get('neededBy'):
        head += f" · by {d['neededBy']}"
    if d.get('from'):
        head += f" · from {d['from']}"
    box = _box(slide.shapes, x, y, w, h, fill='F2F5FE', line=ACCENT, line_pt=1.25, kind=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06, name='Decision needed')
    _write(_frame(box, MSO_ANCHOR.MIDDLE, 0.11), [[(head.upper(), {'bold': True, 'color': ACCENT, 'size': size * 0.86})], [(d.get('text') or '', {})]], size, space_after=2)


def _milestone_table(slide, x, y, w, milestones, size, with_base=False):
    rows = milestones[:8]
    ncol = 6 if with_base else 5
    row_h = size * 1.85 / 72
    shape = slide.shapes.add_table(len(rows) + 1, ncol, Inches(x), Inches(y), Inches(w), Inches(row_h * (len(rows) + 1)))
    shape.name = 'Milestones'
    tbl = shape.table
    style = shape._element.graphic.graphicData.tbl.tblPr.find(qn('a:tableStyleId'))
    if style is not None:
        style.text = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'          # "No Style, No Grid": ours is drawn below
    tbl.first_row = True
    tbl.horz_banding = False
    for ci, frac in enumerate((0.44, 0.11, 0.11, 0.11, 0.09, 0.14) if with_base else (0.50, 0.17, 0.11, 0.09, 0.13)):
        tbl.columns[ci].width = Inches(w * frac)
    label = {'done': ('●  Met', OK), 'late': ('■  Past due', BAD), 'upcoming': ('◆  Ahead', INK)}
    if with_base:                                  # the audit trail: baseline, forecast, slip
        data = [(('Milestone', 'Baseline', 'Forecast', 'Slip', 'Done', 'Status'), True)] + [
            ((m['title'], _fmt_day(m['baseline']) if m.get('baseline') else 'new', _fmt_day(m['date']), _slip_text(m.get('slip_days')),
              f"{m.get('percent_complete', 0)}%", label[m['state']][0]), False) for m in rows]
    else:
        data = [(('Milestone', 'Track', 'Date', 'Done', 'Status'), True)] + [
            ((m['title'], m.get('category') or '', _fmt_day(m['date']), f"{m.get('percent_complete', 0)}%", label[m['state']][0]), False) for m in rows]
    last = ncol - 1
    for ri, (cells, head) in enumerate(data):
        tbl.rows[ri].height = Inches(row_h)
        for ci, val in enumerate(cells):
            cell = tbl.cell(ri, ci)
            _cell_lines(cell, bottom=RULE, weight=12700 if head else 6350)
            cell.fill.background()
            cell.margin_left = cell.margin_top = cell.margin_bottom = Inches(0.0)
            cell.margin_right = Inches(0.08)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            late = with_base and ci == 3 and not head and (rows[ri - 1].get('slip_days') or 0) > 0
            color = INK3 if head else (label[rows[ri - 1]['state']][1] if ci == last else BAD if late else INK)
            _write(tf, [[(val.upper() if head else val, {'bold': head or ci == last or late})]], size * (0.84 if head else 1), color=color)
    return row_h * (len(rows) + 1)


TREND_MARKS = (MSO_SHAPE.OVAL, MSO_SHAPE.RECTANGLE, MSO_SHAPE.ISOSCELES_TRIANGLE, MSO_SHAPE.DIAMOND, MSO_SHAPE.MATH_MULTIPLY)


def _trend_series(facts, milestones):
    """Drift per series, mirroring MilestoneTrend.jsx: days later than the date first reported."""
    now = {'as_of': facts['as_of'], 'end': facts.get('end'), 'milestones': {str(m['id']): m['end'] for m in milestones}}
    pts = ((facts.get('history') or []) + [now])[-8:]
    if len(pts) < 3:
        return pts, []
    raw = [('Finish', [p.get('end') for p in pts], True)] + [
        (m['title'], [(p.get('milestones') or {}).get(str(m['id'])) for p in pts], False) for m in milestones if m['state'] != 'done'][:4]
    series = []
    for title, values, strong in raw:
        if sum(1 for v in values if v) < 2:
            continue
        first = next(v for v in values if v)
        d0 = datetime.fromisoformat(first).timestamp() // 86400
        series.append({'title': title, 'strong': strong, 'values': values,
                       'drift': [None if not v else int(round(datetime.fromisoformat(v).timestamp() / 86400 - d0)) for v in values]})
    return pts, series


def _draw_trend(slide, x0, y0, w, h, facts, milestones, tz, pt):
    """The milestone trend chart as grouped native lines and markers."""
    pts, series = _trend_series(facts, milestones)
    if not series:
        return None
    grp = slide.shapes.add_group_shape()
    g = grp.shapes
    u = pt / 72
    _text(g, x0, y0, w, u * 1.5, [[('MILESTONE TREND · DAYS LATER THAN FIRST REPORTED', {'bold': True})]], pt * 0.88, color=INK3, name='Trend title')
    left, right, top, bottom = x0 + u * 4.2, x0 + w - w * 0.33, y0 + u * 2.2, y0 + h - u * 1.7
    vals = [d for s_ in series for d in s_['drift'] if d is not None]
    lo, hi = min(0, *vals), max(0, *vals)
    if hi - lo < 4:
        hi += math.ceil((4 - (hi - lo)) / 2)
        lo = min(lo, hi - 4)

    def PX(i):
        return left + (i / (len(pts) - 1)) * (right - left)

    def PY(d):
        return top + (1 - (d - lo) / (hi - lo)) * (bottom - top)

    step = max(1, math.ceil((hi - lo) / 4))
    d = math.ceil(lo / step) * step
    while d <= hi:
        _box(g, left, PY(d) - 0.004, right - left, 0.008, fill=INK3 if d == 0 else RULE, name='Grid')
        _text(g, x0, PY(d) - u * 0.7, u * 3.8, u * 1.4, ['on plan' if d == 0 else _signed(d)], pt * 0.82, color=INK3, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE, name='Axis')
        d += step
    for i, p in enumerate(pts):
        label = 'now' if i == len(pts) - 1 else _fmt_day(_parse(p['as_of'], tz))
        _text(g, PX(i) - u * 3, bottom + u * 0.3, u * 6, u * 1.3, [label], pt * 0.82, color=INK3, align=PP_ALIGN.CENTER, name='Report date')

    ends = sorted(((PY(next(v for v in reversed(s_['drift']) if v is not None)), k) for k, s_ in enumerate(series)))
    label_y, prev = {}, None
    for yy, k in ends:                              # right-hand labels, nudged apart
        yy = yy if prev is None or yy - prev >= u * 1.25 else prev + u * 1.25
        label_y[k], prev = yy, yy
    for k, s_ in enumerate(series):
        color = ACCENT if s_['strong'] else INK
        xy = [(PX(i), PY(v)) for i, v in enumerate(s_['drift']) if v is not None]
        for (ax, ay), (bx, by) in zip(xy, xy[1:]):
            ln = g.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(ax), Inches(ay), Inches(bx), Inches(by))
            ln.line.color.rgb = _rgb(color)
            ln.line.width = Pt(2 if s_['strong'] else 1.1)
            ln.name = f'Trend: {s_["title"]}'
        q = u * 0.3
        for px, py in xy:
            _box(g, px - q, py - q, q * 2, q * 2, fill=color, kind=TREND_MARKS[k % len(TREND_MARKS)], name='Trend point')
        last_v = next(v for v in reversed(s_['values']) if v)
        last_d = next(v for v in reversed(s_['drift']) if v is not None)
        title = s_['title'] if len(s_['title']) <= 22 else s_['title'][:21] + '…'
        text = f"{title} · {_fmt_day(_parse(last_v, tz))}" + (f' ({_signed(last_d)})' if last_d else '')
        _text(g, right + u * 0.8, label_y[k] - u * 0.7, x0 + w - right - u * 0.8, u * 1.4, [[(text, {'bold': s_['strong']})]], pt * 0.9, color=color, anchor=MSO_ANCHOR.MIDDLE, name=f'Trend label: {s_["title"]}')
    _describe(grp, 'Milestone trend', 'Milestone trend: ' + '; '.join(
        f"{s_['title']} {_slip_text(next(v for v in reversed(s_['drift']) if v is not None))}" for s_ in series) + f' across {len(pts)} reports.')
    return grp


def _cell_lines(cell, bottom, weight):
    """python-pptx has no cell-border API. Order matters in <a:tcPr>: lnL, lnR, lnT, lnB, then fill."""
    tc_pr = cell._tc.get_or_add_tcPr()
    for tag in ('a:lnL', 'a:lnR', 'a:lnT', 'a:lnB'):
        for old in tc_pr.findall(qn(tag)):
            tc_pr.remove(old)
    for pos, tag in enumerate(('a:lnL', 'a:lnR', 'a:lnT', 'a:lnB')):
        ln = etree.SubElement(tc_pr, qn(tag))
        tc_pr.remove(ln)
        tc_pr.insert(pos, ln)
        if tag == 'a:lnB':
            ln.set('w', str(weight))
            etree.SubElement(etree.SubElement(ln, qn('a:solidFill')), qn('a:srgbClr')).set('val', bottom)
        else:
            ln.set('w', '0')
            etree.SubElement(ln, qn('a:noFill'))


# ── the two layouts ─────────────────────────────────────────────────────────────────────────────
def build_pptx(*, doc, report, facts, layout='slide', paper='letter', previous=None, tz_offset_minutes=0):
    """Return the .pptx file as bytes."""
    tz = dt_tz(timedelta(minutes=max(-840, min(840, int(tz_offset_minutes or 0)))))
    status = STATUS.get(report.get('status'), STATUS['on_track'])
    report = {**report, 'status': report.get('status') if report.get('status') in STATUS else 'on_track'}
    show = {'pathToGreen': True, 'decision': True, 'kpis': True, 'timeline': True, 'baseline': False, 'moved': False, 'trend': False,
            'columns': True, 'milestoneTable': True, 'footer': True, **(doc.get('show') or {})}
    handout = layout == 'handout'

    prs = Presentation()
    pw, ph = PAPER.get(paper, PAPER['letter']) if handout else (13.333, 7.5)
    prs.slide_width, prs.slide_height = Inches(pw), Inches(ph)
    slide = prs.slides.add_slide(prs.slide_layouts[6])                      # blank layout
    sh = slide.shapes

    if handout:
        L, R, T, B, GAP = 0.51, 0.51, 0.43, 0.31, 0.14
        S = {'ident': 11, 'chip': 10.5, 'headline': 21, 'ptg': 11, 'ask': 11, 'kpi': (8.5, 16.5, 9.5), 'tl': 10, 'block': (8.8, 10.5), 'table': 10.5, 'foot': 8.5}
    else:
        L, R, T, B, GAP = 0.40, 0.40, 0.30, 0.21, 0.14
        S = {'ident': 11, 'chip': 10.5, 'headline': 20, 'ptg': 11, 'ask': 11.5, 'kpi': (10, 17.5, 10), 'tl': 10, 'block': (10, 11), 'table': 10.5, 'foot': 10}
    cw = pw - L - R
    header = doc.get('header') or {}
    y = T

    # 1. identity + status chip
    ident = [(header.get('project') or '', {'bold': True, 'color': INK})]
    for part in (header.get('subtitle'), header.get('date'), header.get('pm')):
        if part:
            ident += [('   |   ', {'color': RULE}), (part, {})]
    chip_text = f"{status['glyph']}  {status['label'].upper()}"
    trend = _trend(report, previous, tz)
    chip_w = min(cw * 0.5, _est_width(chip_text, S['chip'], True) + _est_width(trend, S['chip'] * 0.95) + 0.45)
    _text(sh, L, y, cw - chip_w - 0.2, 0.32, [ident], S['ident'], color=INK2, anchor=MSO_ANCHOR.MIDDLE, name='Report identity')
    chip = _box(sh, L + cw - chip_w, y, chip_w, 0.32, fill=status['fill'], line=status['color'], line_pt=1, kind=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5, name='Status')
    _write(_frame(chip, MSO_ANCHOR.MIDDLE, 0.06), [[(chip_text, {'bold': True, 'color': status['color']}), (trend, {'color': INK2, 'size': S['chip'] * 0.95})]], S['chip'], align=PP_ALIGN.CENTER)
    _describe(chip, 'Status', f"Status: {status['label']}{trend}")
    y += 0.32 + GAP

    # 2. headline (+ path to green) and the ask
    ptg = (doc.get('pathToGreen') or '').strip() if (show['pathToGreen'] and report['status'] != 'on_track') else ''
    ask_on = bool(show['decision'])
    head_w = cw * 0.62 if (ask_on and not handout) else cw
    head_lines = min(3, _lines(doc.get('headline') or '', S['headline'], head_w, True)) or 1
    ptg_lines = min(3, _lines('Path to green: ' + ptg, S['ptg'], head_w)) if ptg else 0
    lead_h = head_lines * S['headline'] * 1.16 / 72 + (0.09 + ptg_lines * S['ptg'] * 1.3 / 72 if ptg else 0) + 0.06
    if ask_on and not handout:
        lead_h = max(lead_h, 0.78)                  # the ask box beside it needs room for two lines
    paras = [[(doc.get('headline') or '', {'bold': True, 'font': HEAD_FONT, 'size': S['headline']})]]
    if ptg:
        paras.append([('Path to green: ', {'bold': True, 'color': INK, 'size': S['ptg']}), (ptg, {'color': INK2, 'size': S['ptg']})])
    _text(sh, L, y, head_w, lead_h, paras, S['headline'], space_after=5, line_spacing=0.95, name='Headline')
    if ask_on and not handout:
        _ask(slide, L + cw * 0.648, y, cw * 0.352, lead_h, doc, S['ask'])
    y += lead_h + GAP
    if ask_on and handout:
        d_ = doc.get('decision') or {}
        ah = 0.4 if d_.get('none') else 0.34 + _lines(d_.get('text') or '', S['ask'], cw - 0.25) * S['ask'] * 1.3 / 72
        _ask(slide, L, y, cw, ah, doc, S['ask'])
        y += ah + GAP

    # Heights of everything below, so the timeline can take what is left (as on the page).
    kpis = (doc.get('kpis') or [])[:6] if show['kpis'] else []
    per_row = (3 if len(kpis) > 4 else max(1, len(kpis))) if handout else max(1, len(kpis))
    kpi_h = (0.68 if handout else 0.88) * math.ceil(len(kpis) / per_row) if kpis else 0
    cols = [c for c in (doc.get('columns') or []) if not c.get('hidden')] if show['columns'] else []
    milestones = _chosen_milestones(doc, facts, tz) if not facts.get('empty') else []
    table_on = handout and show['milestoneTable'] and milestones
    table_h = (S['table'] * 1.85 / 72) * (min(8, len(milestones)) + 1) if table_on else 0
    foot_lines = min(2, _lines(_footer_text(doc, report, facts, tz), S['foot'], cw - 0.9)) or 1      # a baseline makes it longer
    foot_h = (0.10 + foot_lines * S['foot'] * 1.3 / 72) if show['footer'] else 0
    with_base = bool(facts.get('baseline')) and show['baseline'] is True
    moved_text = (doc.get('moved') or '').strip() if (show['moved'] is True and facts.get('since_last')) else ''
    moved_lead = f"Moved since {_fmt_day(_parse(facts['since_last']['as_of'], tz))}: " if moved_text else ''
    moved_h = _lines(moved_lead + moved_text, S['ptg'], cw) * S['ptg'] * 1.3 / 72 + 0.04 if moved_text else 0
    trend_on = handout and show['trend'] is True and bool(_trend_series(facts, milestones)[1])
    trend_h = 1.45 if trend_on else 0

    def block_h(c):
        # Estimated from the wording and this block's width, so the space reserved is the space used.
        bw = (cw if handout and len(cols) == 1 else (cw - 0.29) / 2) if handout else (cw - 0.29 * max(0, len(cols) - 1)) / max(1, len(cols))
        if c.get('kind') == 'text':
            lines = max(1, _lines(c.get('text') or '', S['block'][1], bw))
        else:
            items = [i for i in c.get('items') or [] if i.get('text') or i.get('detail')]
            lines = sum(_lines(f"{i.get('text') or ''} {i.get('detail') or i.get('when') or ''}", S['block'][1], bw - 0.2, c.get('kind') == 'risks') + 0.28 for i in items) or 1
        return 0.36 + lines * S['block'][1] * 1.24 / 72
    if handout:
        pairs = [cols[i:i + 2] for i in range(0, len(cols), 2)]
        cols_h = sum(max(block_h(c) for c in pr) for pr in pairs) + GAP * max(0, len(pairs) - 1)
    else:
        cols_h = max((block_h(c) for c in cols), default=0)
    timeline_on = show['timeline'] and not facts.get('empty') and facts.get('rows')
    fixed = kpi_h + cols_h + table_h + foot_h + moved_h + trend_h + GAP * sum(1 for v in (kpi_h, cols_h, table_h, foot_h, moved_h, trend_h) if v)
    avail = ph - B - y - fixed - (GAP if timeline_on else 0)
    # The timeline takes what is left, but never less than a readable height: if the page is
    # overfull, the blocks below run long rather than the chart collapsing (the print tool warns).
    n_rows = len([r for r in facts.get('rows') or [] if r['name'] not in set((doc.get('timeline') or {}).get('hiddenRows') or [])])
    need = 0.62 + 0.22 * max(1, n_rows)
    tl_h = (max(need, min(avail, 2.1 if handout else 9)) if timeline_on else 0)

    # 3. numbers
    if kpis:
        _kpis(slide, L, y, cw, kpi_h, kpis, S['kpi'], per_row)
        y += kpi_h + GAP
    # 4. timeline
    if timeline_on and tl_h >= 0.6:
        _draw_timeline(slide, L, y, cw, tl_h, doc, facts, tz, S['tl'])
        y += tl_h + GAP
    if moved_text:
        _text(sh, L, y, cw, moved_h, [[(moved_lead, {'bold': True, 'color': INK}), (moved_text, {'color': INK2})]], S['ptg'], name='Moved since last report')
        y += moved_h + GAP
    # 5. milestone table (handout), then the trend chart
    if table_on:
        y += _milestone_table(slide, L, y, cw, milestones, S['table'], with_base) + GAP
    if trend_on:
        _draw_trend(slide, L, y, cw, trend_h, facts, milestones, tz, S['tl'])
        y += trend_h + GAP
    # 6. text blocks
    if cols:
        if handout:
            for pr in pairs:
                hh = max(block_h(c) for c in pr)
                bw = cw if len(pr) == 1 else (cw - 0.29) / 2
                for k, c in enumerate(pr):
                    _block(slide, L + k * (bw + 0.29), y, bw, hh, c, S['block'])
                y += hh + GAP
        else:
            weights = [1.45 if c.get('kind') == 'risks' else 1.0 for c in cols]
            gap, total = 0.29, sum(weights)
            usable = cw - gap * (len(cols) - 1)
            cx = L
            for c, wt in zip(cols, weights):
                bw = usable * wt / total
                _block(slide, cx, y, bw, cols_h, c, S['block'])
                cx += bw + gap
            y += cols_h + GAP
    # 7. footer
    if show['footer']:
        fy = ph - B - foot_h
        _box(sh, L, fy, cw, 0.012, fill=RULE, name='Rule')
        _text(sh, L, fy + 0.06, cw - 0.9, foot_h - 0.06, [_footer_text(doc, report, facts, tz)], S['foot'], color=INK3, name='Footer')
        _text(sh, L + cw - 0.9, fy + 0.06, 0.9, 0.2, [settings.APP_NAME], S['foot'], color=INK3, align=PP_ALIGN.RIGHT, name='Source')

    # Speaker notes: provenance travels with the slide even if the footer is deleted.
    notes = slide.notes_slide.notes_text_frame
    notes.text = (f"{header.get('project') or ''} — status report. {_footer_text({'footer': {}}, report, facts, tz)}. "
                  f'Generated by {settings.APP_NAME}; every element on this slide is a native, editable PowerPoint object.')
    prs.core_properties.title = f"{header.get('project') or 'Project'} — status report"
    prs.core_properties.subject = STATUS[report['status']]['label']

    out = BytesIO()
    prs.save(out)
    return out.getvalue()
